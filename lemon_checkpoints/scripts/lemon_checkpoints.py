#!/usr/bin/env python3
"""Lemon Checkpoints v1.2：生成输入检查、执行底稿、Excel（默认）、脑图、Markdown或提示词。"""

from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def config() -> dict:
    return json.loads((ROOT / "config.json").read_text(encoding="utf-8"))


def source_text(path: str) -> str:
    item = Path(path)
    if not item.is_file():
        raise SystemExit(f"找不到输入文件：{item}")
    text = item.read_text(encoding="utf-8").strip()
    if not text:
        raise SystemExit("输入文件没有内容。")
    return text


def found(text: str, terms: list[str]) -> list[str]:
    lowered = text.casefold()
    return [term for term in terms if term.casefold() in lowered]


def input_mode(text: str, settings: dict) -> str:
    size = len(re.sub(r"\s+", "", text))
    structured = bool(re.search(r"(^|\n)#{1,4}\s|(^|\n)\s*[-*]\s|背景|流程|规则|限制", text))
    return "补充信息" if size < settings["short_input_limit"] and not structured else "直接整理"


def signals(text: str, settings: dict) -> dict:
    return {
        "含糊表达": found(text, settings["uncertain_terms"]),
        "计算类场景": found(text, settings["calculation_terms"]),
        "智能功能": found(text, settings["intelligence_terms"]),
    }


def scan_markdown(text: str, settings: dict) -> str:
    detected = signals(text, settings)
    lines = [
        "# 输入检查",
        "",
        f"检查日期：{date.today().isoformat()}",
        f"处理方式：{input_mode(text, settings)}",
        "",
        "## 材料判断",
        "",
        "| 项目 | 当前判断 | 后续动作 |",
        "|---|---|---|",
        "| 业务背景 | [从材料提取] | 确认本次测试对象和边界 |",
        "| 已确认事实 | [从原文或截图提取] | 编号为F-XXX |",
        "| 验证建议 | [结合风险提出] | 不写成既定规则 |",
        "| 待确认问题 | [识别缺失和冲突] | 编号为O-XXX |",
        "",
        "## 风险信号",
        "",
    ]
    for label, values in detected.items():
        lines.append(f"- {label}：{', '.join(values) if values else '未识别'}")
    lines.extend(["", "> 关键词只负责提醒，不代表实际覆盖率。", ""])
    return "\n".join(lines)


def evidence_rows() -> list[list[str]]:
    return [
        ["F-001", "已确认", "[从材料直接提取]", "[原文段落、截图区域或接口文档]", "有效"],
        ["V-001", "验证建议", "[高风险但材料未说明的检查]", "[关联O-XXX]", "待评估"],
    ]


def question_rows() -> list[list[str]]:
    return [[
        "O-001", "[会改变步骤、预期结果或风险判断的问题]", "[关联LC-XXX]",
        "[规则缺失造成的测试或业务风险]", "", "开放", "",
    ]]


def checkpoint_rows(settings: dict) -> list[list[str]]:
    batches = [
        ("A", "首轮冒烟", "S0", "功能"),
        ("B", "重点验证", "S1", "安全"),
        ("C", "扩展检查", "S2", "兼容"),
    ]
    rows = []
    number = 1
    for batch, batch_name, severity, risk in batches:
        for route in settings["routes"]:
            rows.append([
                f"LC-{number:03d}", f"{batch} {batch_name}", route, risk,
                "[具体触发情况]", "[账号、数值、状态或组合]", "[页面、接口、数据或日志中的可核对结果]",
                "验证建议", "V-001 / O-001", severity, "[人工/自动化/组合 + 原因]", "[直接业务影响]",
            ])
            number += 1
    return rows


def coverage_rows(settings: dict) -> list[list[str]]:
    return [[code, name, "待补充", "[关联LC-XXX]", "[缺少规则、数据或环境]"] for code, name in settings["coverage"]]


def markdown_build(text: str, settings: dict) -> str:
    lines = [
        "# 测试执行清单",
        "",
        f"生成日期：{date.today().isoformat()}",
        f"处理方式：{input_mode(text, settings)}",
        "",
        "## 背景和材料边界",
        "",
        "[根据输入说明业务目标、测试对象和不可见信息。]",
        "",
        "## 输入底稿",
        "",
        "| 编号 | 证据级别 | 内容 | 来源或影响范围 | 状态 |",
        "|---|---|---|---|---|",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in evidence_rows())
    lines.extend([
        "", "## 待确认清单", "",
        "| 编号 | 需要确认的问题 | 影响测试点 | 风险或影响 | 负责人 | 状态 | 确认结论 |",
        "|---|---|---|---|---|---|---|",
    ])
    lines.extend("| " + " | ".join(row) + " |" for row in question_rows())
    lines.extend([
        "", "## 风险路线", "",
        "进入与展示 → 信息输入 → 系统校验 → 结果反馈 → 状态保持与退出",
        "", "## 执行清单", "",
        "| 检查编号 | 执行批次 | 业务链路 | 风险主题 | 触发情况 | 样本数据 | 核对结果 | 证据级别 | 关联项 | 严重度 | 执行方式 | 失败影响 |",
        "|---|---|---|---|---|---|---|---|---|---|---|---|",
    ])
    lines.extend("| " + " | ".join(row) + " |" for row in checkpoint_rows(settings))
    lines.extend([
        "", "## 覆盖地图", "",
        "| 代码 | 检查方向 | 状态 | 关联测试点 | 当前缺口 |",
        "|---|---|---|---|---|",
    ])
    lines.extend("| " + " | ".join(row) + " |" for row in coverage_rows(settings))
    return "\n".join(lines) + "\n"


def mindmap_build(settings: dict) -> str:
    grouped = {"A 首轮冒烟": [], "B 重点验证": [], "C 扩展检查": []}
    for row in checkpoint_rows(settings):
        grouped[row[1]].append(row)
    lines = ["# [需求名称]测试执行图", "", "- 输入底稿", "  - F-001 已确认｜[材料事实]", "  - O-001 待确认｜[开放问题]", "  - V-001 验证建议｜[风险检查]", ""]
    for batch, rows in grouped.items():
        lines.append(f"- {batch}")
        for row in rows:
            lines.extend([f"  - {row[0]} [{row[9]}] {row[2]}｜{row[4]}", f"    - 数据：{row[5]}", f"    - 结果：{row[6]}", f"    - 依据：{row[8]}"])
        lines.append("")
    lines.append("- 覆盖地图")
    for row in coverage_rows(settings):
        lines.append(f"  - {row[0]} {row[1]}｜{row[2]}｜{row[3]}")
    return "\n".join(lines) + "\n"


def excel_build(text: str, settings: dict, output: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise SystemExit("Excel输出需要 openpyxl。请先运行：pip install -r requirements.txt") from exc
    destination = Path(output)
    if destination.suffix.lower() != ".xlsx":
        raise SystemExit("Excel输出文件名必须以 .xlsx 结尾。")
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "执行看板"
    points = workbook.create_sheet("测试清单")
    questions = workbook.create_sheet("待确认")
    evidence = workbook.create_sheet("规则证据")
    coverage = workbook.create_sheet("覆盖地图")

    point_headers = ["检查编号", "执行批次", "业务链路", "风险主题", "触发情况", "样本数据", "核对结果", "证据级别", "关联项", "严重度", "执行方式", "失败影响"]
    points.append(point_headers)
    for row in checkpoint_rows(settings): points.append(row)
    questions.append(["编号", "需要确认的问题", "影响测试点", "风险或影响", "负责人", "状态", "确认结论"])
    for row in question_rows(): questions.append(row)
    evidence.append(["编号", "证据级别", "内容", "来源或影响范围", "状态"])
    for row in evidence_rows(): evidence.append(row)
    coverage.append(["代码", "检查方向", "状态", "关联测试点", "当前缺口"])
    for row in coverage_rows(settings): coverage.append(row)
    dashboard.append(["测试执行看板", "数量"])
    dashboard.append(["A 首轮冒烟", "=COUNTIF('测试清单'!$B$2:$B$100,\"A*\")"])
    dashboard.append(["B 重点验证", "=COUNTIF('测试清单'!$B$2:$B$100,\"B*\")"])
    dashboard.append(["C 扩展检查", "=COUNTIF('测试清单'!$B$2:$B$100,\"C*\")"])
    dashboard.append(["S0", "=COUNTIF('测试清单'!$J$2:$J$100,\"S0\")"])
    dashboard.append(["开放问题", "=COUNTIF('待确认'!$F$2:$F$100,\"开放\")"])
    dashboard.append(["材料处理方式", input_mode(text, settings)])

    widths = {
        "执行看板": [24, 24],
        "测试清单": [12, 16, 20, 14, 34, 28, 40, 14, 18, 10, 28, 30],
        "待确认": [12, 48, 24, 34, 14, 12, 32],
        "规则证据": [12, 14, 48, 32, 12],
        "覆盖地图": [10, 20, 14, 26, 38],
    }
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="365314")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
        for index, width in enumerate(widths[sheet.title], 1): sheet.column_dimensions[get_column_letter(index)].width = width
    for value, color in (("S0", "FECACA"), ("S1", "FED7AA"), ("S2", "FEF3C7")):
        points.conditional_formatting.add(f"J2:J{points.max_row}", FormulaRule(formula=[f'J2="{value}"'], fill=PatternFill("solid", fgColor=color)))
    workbook.save(destination)
    print(f"已生成：{destination}")


def prompt_build(text: str) -> str:
    template = (ROOT / "assets" / "prompt-template.md").read_text(encoding="utf-8")
    return template.replace("[粘贴需求、截图说明、用户故事或接口材料]", text) + "\n"


def write_text(content: str, output: str | None) -> None:
    if output:
        target = Path(output)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8")
        print(f"已生成：{target}")
    else:
        print(content, end="")


def arguments() -> argparse.Namespace:
    settings = config()
    default_format = settings.get("default_format", "excel")
    if default_format not in {"markdown", "mindmap", "excel"}:
        default_format = "excel"
    parser = argparse.ArgumentParser(description="Lemon Checkpoints v1.2")
    commands = parser.add_subparsers(dest="command", required=True)
    for name in ("scan", "inspect", "build", "draft", "prompt"):
        item = commands.add_parser(name)
        item.add_argument("-i", "--input", required=True)
        item.add_argument("-o", "--output")
        if name in ("build", "draft"):
            item.add_argument(
                "--format",
                choices=["markdown", "mindmap", "excel"],
                default=default_format,
            )
            item.add_argument("--layout", choices=["table", "mindmap", "excel"])
    return parser.parse_args()


def default_excel_output(input_path: str) -> str:
    """未指定 -o 时，按输入文件名生成同目录 *-测试执行清单.xlsx。"""
    source = Path(input_path)
    return str(source.with_name(f"{source.stem}-测试执行清单.xlsx"))


def main() -> None:
    args = arguments()
    settings = config()
    text = source_text(args.input)
    command = {"inspect": "scan", "draft": "build"}.get(args.command, args.command)
    if command == "scan":
        write_text(scan_markdown(text, settings), args.output)
    elif command == "prompt":
        write_text(prompt_build(text), args.output)
    else:
        requested = args.layout or args.format
        fmt = "markdown" if requested == "table" else requested
        if fmt == "excel":
            output = args.output or default_excel_output(args.input)
            excel_build(text, settings, output)
        elif fmt == "mindmap":
            write_text(mindmap_build(settings), args.output)
        else:
            write_text(markdown_build(text, settings), args.output)


if __name__ == "__main__":
    main()
